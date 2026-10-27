"""Turn cluster_E_timetable_dataset.xlsx (5th semester) into neo4j/02_cluster_E_sem5.cypher.

Usage:  python build_cluster_E_sem5_cypher.py "C:\\path\\to\\cluster_E_timetable_dataset.xlsx"

The workbook has three sheets: Sessions, Subjects_Faculty, Metadata. Its sessions use the bell
schedule of time grid A (tea 10:45-11:15, lunch 14:00-14:55).
"""
import os
import re
import sys
from collections import OrderedDict

import openpyxl

from timetable_common import Builder, write_cypher, to24

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\dell\Downloads\cluster_E_timetable_dataset.xlsx"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "02_cluster_E_sem5.cypher")

DEPT, PROGRAM, YEAR, SEM, CLUSTER, GRID = "CSE", "UG", "2026-27", 5, "E", "A"
SEM_ID = f"{DEPT}-{PROGRAM}-{YEAR}-S{SEM}"
b = Builder()

# ------------------------------------------------------------------ read the workbook
wb = openpyxl.load_workbook(SRC, data_only=True)
meta = {r[0]: r[1] for r in wb["Metadata"].iter_rows(min_row=2, values_only=True) if r[0]}
sessions = [r for r in wb["Sessions"].iter_rows(min_row=2, values_only=True) if r[0]]
subjects = [r for r in wb["Subjects_Faculty"].iter_rows(min_row=2, values_only=True) if r[0]]

wef_text = meta["Effective from"]  # dd/mm/yyyy
d, m, y = wef_text.split("/")
WEF = f"{y}-{m}-{d}"
SECTIONS = OrderedDict(re.findall(r"(\w+) \[(\w+)\]", meta["Classes included"]))   # 5E1 -> 5A, ...
sec_id = lambda s: f"{SEM_ID}-{s}"
fac_id = lambda ini: f"{DEPT}-{ini}"

# ------------------------------------------------------------------ courses and faculty
SHORT = {"AE": "CC", "Mini Project": "MIP"}                       # session label -> short name
FRIENDLY = {"PE": "Professional Elective", "AE": "Ability Enhancement Course"}
courses = OrderedDict()            # code -> dict
teacher = {}                       # (section, subject) -> initials or None
faculty_names = {}                 # initials -> full name
for sec, subj, code, name, fname, ini, ltp in subjects:
    c = courses.setdefault(code, dict(code=code, subject=subj, name=FRIENDLY.get(subj, name), short=SHORT.get(subj, subj),
                                      ltp=ltp, scheduling="REGULAR"))
    if c["ltp"] != ltp:
        sys.exit(f"LTP differs for {code}: {c['ltp']} vs {ltp}")
    if code.startswith("23CS5PC") or code.startswith("23CS5BS"):
        cat = "PC" if "PC" in code[:8] else "BS"
    elif code.startswith("23CV5HS"):
        cat = "HS"
    elif "SPE" in code:
        cat = "PE"
    elif "SAE" in code:
        cat = "AE"
    elif "PWMIP" in code:
        cat = "PW"
    else:
        cat = "PC"
    c["category"] = cat
    teacher[(sec, subj)] = ini
    if ini and fname and fname != "Not specified":
        faculty_names.setdefault(ini, fname)
courses_by_subject = {c["subject"]: c for c in courses.values()}
for c in courses.values():
    if c["subject"] == "PE":
        c["scheduling"] = "ELECTIVE_SLOT"
    if c["subject"] == "Mini Project":
        c["scheduling"] = "NONE"
b.note("PE is an elective slot (reserves the period only) and Mini Project is not timetabled; both have no teacher in the file.")
b.note("The sessions called CC LAB are taken as the lab of the AE course (same assumption as in the 3rd semester file).")
b.note("Course codes follow this file: PE = 23CS5SPEXXX, AE = 23CS5SAEXXX (the earlier printed timetable showed 23CS5PEXXX / 23CS5AEXXX). "
       "The clean-up block removes the old two courses so they are not allotted twice.")

FACULTY = OrderedDict((fac_id(i), dict(initials=i, name=n, dept=DEPT, cluster=True)) for i, n in faculty_names.items())


def add_faculty(ini):
    FACULTY.setdefault(fac_id(ini), dict(initials=ini, name=None, dept=DEPT, cluster=True))


BATCHES = []
CLASSES = OrderedDict()
ALLOTS = []
for sec in SECTIONS:
    for code, c in courses.items():
        t = teacher.get((sec, c["subject"]))
        ALLOTS.append(dict(id=f"{sec_id(sec)}-{code}", section=sec_id(sec), course=code, teacher=fac_id(t) if t else None))

lab_batches = {}          # (section, subject) -> [(batch, initials)]
for sec, day, start, end, subj, room, stype, batch_details in sessions:
    if batch_details:
        label = subj[:-4] if subj.endswith(" LAB") else subj
        subject = "AE" if label == "CC" else label
        pairs = [tuple(x.strip().split("-", 1)) for x in batch_details.split(";")]
        lab_batches[(sec, subject)] = pairs
        for _, ini in pairs:
            add_faculty(ini)

for a in ALLOTS:
    sec = a["section"].rsplit("-", 1)[1]
    c = courses[a["course"]]
    for batch, ini in lab_batches.get((sec, c["subject"]), []):
        BATCHES.append(dict(id=f"{a['id']}-{batch}", allot=a["id"], name=batch, faculty=fac_id(ini), type="LAB"))
for t in list(teacher.values()):
    if t:
        add_faculty(t)

for sec, day, start, end, subj, room, stype, batch_details in sessions:
    slots = b.slots(GRID, day, to24(start), to24(end), subj)
    if subj.lower().startswith("seminar"):
        item = dict(id=f"{sec_id(sec)}-A-{slots[0]}", name=subj, section=sec_id(sec), allot=None, type="A",
                    slots=slots, rooms=[], faculty=[], batches=[])
    else:
        if subj.endswith(" LAB"):
            label, typ = subj[:-4], "P"
        elif subj.endswith("_T"):
            label, typ = subj[:-2], "T"
        else:
            label, typ = subj, "L"
        subject = "AE" if label == "CC" else label
        course = courses_by_subject[subject]
        allot = f"{sec_id(sec)}-{course['code']}"
        rooms = b.room_ids(room) if room else []
        if typ == "P":
            bs = [x for x in BATCHES if x["allot"] == allot]
            fac = list(OrderedDict((x["faculty"], 1) for x in bs))
            batches = [x["id"] for x in bs]
        else:
            t = teacher.get((sec, subject))
            fac, batches = ([fac_id(t)] if t else []), []
        name = {"L": course["short"], "T": course["short"] + " Tutorial", "P": course["short"] + " Lab"}[typ]
        item = dict(id=f"{allot}-{typ}-{slots[0]}", name=name, section=sec_id(sec), allot=allot, type=typ,
                    slots=slots, rooms=rooms, faculty=fac, batches=batches)
    CLASSES.setdefault(sec, []).append(item)

b.note("DEV_T sessions are the DEV tutorial (type T, two periods); the sheet lists them as 'Lecture'.")
b.note("EVS (1-0-0) has no session in the file, so it has no class yet. The allotment in Part C of 01_schema.cypher can place it.")
b.note(f"Effective from {wef_text}; the earlier printed 5E4 timetable said 21/09/2026, so the file's date is used.")
b.note("Saturday is a working day (01_schema.cypher), because 5E1, 5E2 and 5E3 have Saturday classes.")
b.note("Room blocks: all numbered rooms and 'LAB nnn' are put in block PJ on the floor given by the first digit (4F LAB = floor 4). "
       "Only 406/407 were confirmed as PJ earlier; the rest are assumed.")
b.note("These sections follow time grid A: tea 10:45-11:15, lunch 14:00-14:55.")

write_cypher(
    OUT, b,
    header=[f"Cluster {CLUSTER}, semester {SEM}: {', '.join(f'{n} [{c}]' for n, c in SECTIONS.items())}",
            f"Generated from {os.path.basename(SRC)} by tools/build_cluster_E_sem5_cypher.py",
            f"Source: {meta.get('Source', '')}",
            f"{meta.get('College', '')} - {meta.get('Department', '')}",
            f"{meta.get('Timetable period', '')}, effective from {wef_text}"],
    dept=DEPT, program=PROGRAM, year=YEAR, sem=SEM, term="ODD", cluster=CLUSTER, grid=GRID,
    sections=SECTIONS, wef=WEF, courses=courses, faculty=FACULTY, allots=ALLOTS, batches=BATCHES, classes=CLASSES,
    remove_courses=["23CS5PEXXX", "23CS5AEXXX"], remove_sections=["5E5"])
