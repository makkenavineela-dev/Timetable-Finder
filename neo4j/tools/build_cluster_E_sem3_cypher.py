"""Turn Cluster_E_3rd_Sem_Timetable_Extracted.xlsx (3rd semester) into neo4j/02_cluster_E_sem3.cypher.

Usage:  python build_cluster_E_sem3_cypher.py "C:\\path\\to\\Cluster_E_3rd_Sem_Timetable_Extracted.xlsx"

Sheets read: Read Me, All Class Slots (the sessions), Subjects & Credits, Faculty by Class.
(The per-section sheets 3E1..3E4 repeat All Class Slots; they are checked against it.)
The sessions follow bell schedule B: tea 09:50-10:20, lunch 13:05-14:00, eight periods.
"""
import os
import re
import sys
from collections import OrderedDict

import openpyxl

from timetable_common import Builder, write_cypher, to24, DAYS

SRC = sys.argv[1] if len(sys.argv) > 1 else r"C:\Users\dell\Desktop\project k\Cluster_E_3rd_Sem_Timetable_Extracted.xlsx"
OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "02_cluster_E_sem3.cypher")

DEPT, PROGRAM, YEAR, SEM, CLUSTER, GRID = "CSE", "UG", "2026-27", 3, "E", "B"
SEM_ID = f"{DEPT}-{PROGRAM}-{YEAR}-S{SEM}"
MATH = ("MA", "Mathematics")
b = Builder()

# ------------------------------------------------------------------ read the workbook
wb = openpyxl.load_workbook(SRC, data_only=True)
readme = {r[0]: r[1] for r in wb["Read Me"].iter_rows(min_row=2, values_only=True) if r[0]}
sessions = [r for r in wb["All Class Slots"].iter_rows(min_row=2, values_only=True) if r[0]]
subject_rows = [r for r in wb["Subjects & Credits"].iter_rows(min_row=2, values_only=True) if r[0]]
faculty_rows = [r for r in wb["Faculty by Class"].iter_rows(min_row=2, values_only=True) if r[0]]

d, m, y = readme["Effective from"].split("/")
WEF = f"{y}-{m}-{d}"
# 'Classes' = '3E1 (3A), 3E2 (3J), ...'
SECTIONS = OrderedDict(re.findall(r"(\w+) \((\w+)\)", readme["Classes"]))
sec_id = lambda s: f"{SEM_ID}-{s}"

# the per-section sheets must say the same as All Class Slots
for name in SECTIONS:
    own = [tuple(r) for r in wb[name].iter_rows(min_row=2, values_only=True) if r[0] in DAYS]
    allc = [tuple(r[2:7]) for r in sessions if r[0].startswith(name)]
    if own != allc:
        sys.exit(f"Sheet {name} differs from All Class Slots")

# ------------------------------------------------------------------ courses
CATEGORY = [("IP", "IP"), ("PC", "PC"), ("BS", "BS"), ("AE", "AE"), ("SD", "SD"), ("NC", "NC")]   # letters 6-7 of the code
courses = OrderedDict()           # code -> dict
by_label = {}                     # sheet label (DS, DBMS, ...) -> course
for abbr, name, code, ltp in subject_rows:
    cat = next((c for key, c in CATEGORY if code[5:7] == key), None)
    if ltp == "Not specified":
        b.note(f"{abbr} ({code}) has no L-T-P in the file: stored as 0-0-0 and not timetabled (a non-credit course).")
        ltp, sched = "0-0-0", "NONE"
    elif abbr == "Skill Lab-1":
        b.note("Skill Lab-1 has a teacher but no session in the file, so it is stored as not timetabled.")
        sched = "NONE"
    else:
        sched = "REGULAR"
    courses[code] = dict(code=code, name=name, short=abbr, category=cat, ltp=ltp, scheduling=sched)
    by_label[abbr] = courses[code]

# ------------------------------------------------------------------ faculty
FACULTY = OrderedDict()           # faculty id -> dict
teacher = {}                      # (section, subject label) -> faculty id or None


def initials_of(name):
    """Initials for someone the sheet gives no initials for: 'Dr. Shivraj D L' -> SDL."""
    words = re.sub(r"^(Dr|Prof|Ms|Mr|Mrs)\.?\s+", "", name).split()
    return "".join(w[0] for w in words).upper()


for cls, subj, fname, ini in faculty_rows:
    sec = cls.split()[0]
    if ini:
        fid = f"{DEPT}-{ini}"
        FACULTY.setdefault(fid, dict(initials=ini, name=fname, dept=DEPT, cluster=True))
    else:
        ini2 = initials_of(fname)
        fid = f"{MATH[0]}-{ini2}"
        FACULTY.setdefault(fid, dict(initials=ini2, name=fname, dept=MATH[0], cluster=False))
        b.note(f"{fname} (teaches {subj}) has no initials in the file: initials {ini2} are made from the name, and the "
               f"person is placed in the {MATH[1]} department (DMS is a mathematics course), without a cluster.")
    teacher[(sec, subj)] = fid


def add_batch_faculty(ini):
    FACULTY.setdefault(f"{DEPT}-{ini}", dict(initials=ini, name=None, dept=DEPT, cluster=True))


# ------------------------------------------------------------------ allotments and lab batches
ALLOTS = []
for sec in SECTIONS:
    for code, c in courses.items():
        ALLOTS.append(dict(id=f"{sec_id(sec)}-{code}", section=sec_id(sec), course=code, teacher=teacher.get((sec, c["short"]))))

lab_batches = {}          # (section, subject label) -> [(batch, initials)]
for cls, sec, day, time, subj, room, batch_details in sessions:
    if batch_details:
        label = subj[:-4] if subj.endswith(" LAB") else subj
        pairs = [tuple(x.strip().split("-", 1)) for x in batch_details.split(";")]
        lab_batches[(cls.split()[0], label)] = pairs
        for _, ini in pairs:
            add_batch_faculty(ini)

BATCHES = []
for a in ALLOTS:
    sec = a["section"].rsplit("-", 1)[1]
    c = courses[a["course"]]
    for batch, ini in lab_batches.get((sec, c["short"]), []):
        BATCHES.append(dict(id=f"{a['id']}-{batch}", allot=a["id"], name=batch, faculty=f"{DEPT}-{ini}", type="LAB"))

# ------------------------------------------------------------------ sessions -> classes
CLASSES = OrderedDict()
for cls, sec_code, day, time, subj, room, batch_details in sessions:
    sec = cls.split()[0]
    start, end = (to24(x) for x in time.split("-"))
    slots = b.slots(GRID, DAYS[day], start, end, subj)
    # activities without a course: the seminar slot and the diploma-student mathematics bridge course
    if subj.lower().startswith("seminar") or subj.startswith("DIP-MATH"):
        item = dict(id=f"{sec_id(sec)}-A-{slots[0]}", name=subj, section=sec_id(sec), allot=None, type="A",
                    slots=slots, rooms=[], faculty=[], batches=[])
    else:
        if subj.endswith(" LAB"):
            label, typ = subj[:-4], "P"
        elif subj.endswith("_T"):
            label, typ = subj[:-2], "T"
        else:
            label, typ = subj, "L"
        course = by_label[label]
        allot = f"{sec_id(sec)}-{course['code']}"
        rooms = b.room_ids(room) if room else []
        if typ == "P":
            bs = [x for x in BATCHES if x["allot"] == allot]
            fac = list(OrderedDict((x["faculty"], 1) for x in bs))
            batches = [x["id"] for x in bs]
        else:
            t = teacher.get((sec, label))
            fac, batches = ([t] if t else []), []
        name = {"L": course["short"], "T": course["short"] + " Tutorial", "P": course["short"] + " Lab"}[typ]
        item = dict(id=f"{allot}-{typ}-{slots[0]}", name=name, section=sec_id(sec), allot=allot, type=typ,
                    slots=slots, rooms=rooms, faculty=fac, batches=batches)
    CLASSES.setdefault(sec, []).append(item)

b.note("These sections follow time grid B: tea 09:50-10:20, lunch 13:05-14:00, eight periods (08:00 to 16:45). "
       "The file's times were written without am/pm; hours below 8 are read as afternoon.")
b.note("DMS_T sessions are the DMS tutorial (type T, two periods): DMS is 3-1-0.")
b.note("DIP-MATH (A) and DIP-MATH (B) are not in the subject list. They are stored as one-period activities without "
       "course, room or faculty (type A), like the seminar slot; (A) and (B) are kept in the name.")
b.note("The classes that run in several rooms or '4F LAB 1-2' follow the same room rules as the 5th semester file.")
b.note("Faculty names come from 'Faculty by Class'; people who only appear in lab batches have initials only. "
       "Everyone with initials is put in the CSE department and cluster E, as in the 5th semester file.")
b.note(f"Effective from {readme['Effective from']}.")

write_cypher(
    OUT, b,
    header=[f"Cluster {CLUSTER}, semester {SEM}: {', '.join(f'{n} [{c}]' for n, c in SECTIONS.items())}",
            f"Generated from {os.path.basename(SRC)} by tools/build_cluster_E_sem3_cypher.py",
            f"Source: {readme.get('Source', '')}",
            f"{readme.get('College', '')} - {readme.get('Department', '')}",
            f"{readme.get('Period', '')}, effective from {readme['Effective from']}"],
    dept=DEPT, program=PROGRAM, year=YEAR, sem=SEM, term="ODD", cluster=CLUSTER, grid=GRID,
    sections=SECTIONS, wef=WEF, courses=courses, faculty=FACULTY, allots=ALLOTS, batches=BATCHES, classes=CLASSES,
    departments=[MATH])
